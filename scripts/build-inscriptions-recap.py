#!/usr/bin/env python3
"""
Build a filière/région/statut recap workbook from harvested "fiche
d'inscription" PDFs (and, optionally, a dossier-status report) collected
during an authorized engagement to demonstrate the impact of the IDOR
documented in verify-matricules-filieres.js.

Usage:
    python build-inscriptions-recap.py --zip fiches.zip --out recap.xlsx
    python build-inscriptions-recap.py --zip fiches.zip --rapport dossiers.docx --out recap.xlsx

Inputs:
    --zip      Required. ZIP archive of "fiche d'inscription" PDFs (one per
               candidate), as produced by verify-matricules-filieres.js.
    --rapport  Optional. .docx dossier-status report with two tables: one
               listing "OK" matricules (columns: Matricule, Statut, Lien)
               and one listing "PAS OK" matricules (columns: Matricule,
               Existe, Statut, Raison). When given, the workbook gains a
               "Statut dossier" column and two extra cross-tab sheets.
    --out      Required. Output .xlsx path.

Only the tool is meant to be committed to this repo — never the harvested
ZIP/DOCX inputs or the generated .xlsx (see .gitignore): keep evidence
volume and retention to what the engagement's rules of engagement call
for, and delete it once the report is filed.
"""

import argparse
import io
import re
import sys
import zipfile
from collections import Counter

try:
    import pypdf
except ImportError:
    sys.exit("Missing dependency 'pypdf'. Install with: pip install pypdf cryptography cffi")

try:
    import docx
except ImportError:
    docx = None  # only required when --rapport is used

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

# ---------------------------------------------------------------------------
# Tunables — edit here if a future session's fiche wording differs slightly.
# ---------------------------------------------------------------------------
FILIERE_ALIASES = {
    # English-language campus forms use English filière names; harmonize
    # them with the French label used everywhere else so counts aren't split.
    "Gynaecology/Obstetrics": "Gynécologie/Obstétrique",
}
FOREIGN_LABEL = "Étranger (hors Cameroun)"

FIELD_PATTERNS = {
    "MATRICULE": r"MATRICULE\s+(\S+)",
    "NOM": r"NOM\s+(.+?)\s*\n",
    "DATE_NAISSANCE": r"DATE DE NAISSANCE\s+(\S+)",
    "LIEU_NAISSANCE": r"LIEU DE NAISSANCE\s+(.+?)\s*\n",
    "SEXE": r"SEXE\s+(\S+)",
    "TELEPHONE": r"TÉLÉPHONE\s+(\S+)",
    "EMAIL": r"EMAIL\s+(\S+)",
    "NATIONALITE": r"NATIONALITE\s+(.+?)\s+REGION D'ORIGINE",
    "REGION": r"REGION D'ORIGINE\s*(.*?)\s*\n",
    "ETABLISSEMENT": r"ETABLISSEMENT\s*\n?CHOISI\s*\n?(.+?)\nSP[ÉE]CIALIT[ÉE]",
    "SPECIALITE": r"SP[ÉE]CIALIT[ÉE]\s+(.+?)\s*\n",
    "FILIERE": r"FILIERE\s+(.+?)\s*\n",
}
EXAM_TITLE_PATTERN = r"FICHE D'INSCRIPTION\s*:\s*(.+?)\s*MATRICULE"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="1F4E78")
SUB_FONT = Font(name=FONT_NAME, italic=True, size=10, color="595959")
NORMAL_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, bold=True, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ACCENT_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(name=FONT_NAME, size=10, color="9C0006")
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
GREEN_FONT = Font(name=FONT_NAME, size=10, color="006100")

DATA_COLUMNS = [
    ("MATRICULE", "Matricule", 14, "center"),
    ("NOM", "Nom et prénom", 32, None),
    ("SEXE", "Sexe", 11, "center"),
    ("DATE_NAISSANCE", "Date de naissance", 15, "center"),
    ("LIEU_NAISSANCE", "Lieu de naissance", 20, None),
    ("TELEPHONE", "Téléphone", 14, "center"),
    ("EMAIL", "Email", 30, None),
    ("NATIONALITE", "Nationalité", 20, None),
    ("REGION", "Région d'origine", 22, None),
    ("ETABLISSEMENT", "Établissement choisi", 45, None),
    ("FILIERE", "Filière", 30, None),
    ("SPECIALITE", "Spécialité", 30, None),
]
STATUS_COLUMNS = [
    ("STATUT_DOSSIER", "Statut dossier", 15, "center"),
    ("RAISON_STATUT", "Raison (si PAS OK)", 38, None),
]


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------
def parse_fiche_text(text):
    out = {}
    for key, pattern in FIELD_PATTERNS.items():
        m = re.search(pattern, text, re.S if key == "ETABLISSEMENT" else 0)
        out[key] = m.group(1).strip() if m else None

    if out["ETABLISSEMENT"]:
        out["ETABLISSEMENT"] = re.sub(r"\s+", " ", out["ETABLISSEMENT"]).strip()

    if not out["REGION"] or out["REGION"] == "ETABLISSEMENT":
        out["REGION"] = FOREIGN_LABEL
    else:
        out["REGION"] = out["REGION"].upper()

    if out["FILIERE"] in FILIERE_ALIASES:
        out["FILIERE"] = FILIERE_ALIASES[out["FILIERE"]]

    if out["NATIONALITE"]:
        out["NATIONALITE"] = out["NATIONALITE"].replace("é", "É").upper()

    exam_m = re.search(EXAM_TITLE_PATTERN, text, re.S)
    out["_EXAM_TITLE"] = re.sub(r"\s+", " ", exam_m.group(1)).strip() if exam_m else None

    return out


def extract_records_from_zip(zip_path):
    records = []
    skipped = []
    exam_titles = Counter()
    with zipfile.ZipFile(zip_path) as zf:
        pdf_names = [n for n in zf.namelist() if n.lower().endswith(".pdf")]
        for name in sorted(pdf_names):
            data = zf.read(name)
            try:
                reader = pypdf.PdfReader(io.BytesIO(data))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
            except Exception as exc:
                skipped.append((name, f"lecture PDF impossible: {exc}"))
                continue

            rec = parse_fiche_text(text)
            if rec["_EXAM_TITLE"]:
                exam_titles[rec["_EXAM_TITLE"]] += 1
            missing = [k for k in FIELD_PATTERNS if rec[k] is None and k != "REGION"]
            if not rec["MATRICULE"]:
                skipped.append((name, "MATRICULE introuvable dans le texte extrait"))
                continue
            if missing:
                rec["_MISSING_FIELDS"] = missing
            records.append(rec)

    records.sort(key=lambda r: r["MATRICULE"])
    exam_title = exam_titles.most_common(1)[0][0] if exam_titles else None
    return records, skipped, exam_title


# ---------------------------------------------------------------------------
# Optional dossier-status report (.docx)
# ---------------------------------------------------------------------------
def parse_status_report(docx_path):
    if docx is None:
        sys.exit("--rapport requires 'python-docx'. Install with: pip install python-docx")

    document = docx.Document(docx_path)
    status = {}  # matricule -> (statut, raison)

    for table in document.tables:
        if not table.rows:
            continue
        header = [c.text.strip() for c in table.rows[0].cells]
        rows = [[c.text.strip() for c in row.cells] for row in table.rows[1:]]

        if "Raison" in header:
            m_idx = header.index("Matricule")
            r_idx = header.index("Raison")
            for row in rows:
                status[row[m_idx]] = ("PAS OK", row[r_idx])
        elif "Lien" in header or "Statut" in header:
            m_idx = header.index("Matricule")
            for row in rows:
                status.setdefault(row[m_idx], ("OK", ""))

    return status


def merge_status(records, status_map):
    matricules_in_records = {r["MATRICULE"] for r in records}
    for rec in records:
        statut, raison = status_map.get(rec["MATRICULE"], (None, ""))
        rec["STATUT_DOSSIER"] = statut if statut else "Non renseigné"
        rec["RAISON_STATUT"] = raison
    orphan_matricules = sorted(set(status_map) - matricules_in_records)
    return orphan_matricules


# ---------------------------------------------------------------------------
# Workbook construction
# ---------------------------------------------------------------------------
def counts_by(records, field):
    d = Counter(rec.get(field) for rec in records)
    return sorted(d.items(), key=lambda kv: (-kv[1], kv[0] or ""))


def style_header_row(ws, row, headers, widths=None):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def build_workbook(records, out_path, exam_title, orphan_matricules, has_status):
    n = len(records)
    wb = Workbook()

    # --- Données -----------------------------------------------------------
    ws = wb.active
    ws.title = "Données"
    cols = list(DATA_COLUMNS) + (STATUS_COLUMNS if has_status else [])

    style_header_row(ws, 1, [c[1] for c in cols], [c[2] for c in cols])
    ws.row_dimensions[1].height = 28

    for r_idx, rec in enumerate(records, start=2):
        for c_idx, (key, _, _, align) in enumerate(cols, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=rec.get(key))
            cell.font = NORMAL_FONT
            cell.border = BORDER
            if align:
                cell.alignment = Alignment(horizontal=align)

    last_row = n + 1
    table_ref = f"A1:{get_column_letter(len(cols))}{last_row}"
    table = Table(displayName="DonneesCandidats", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"

    col_letter = {key: get_column_letter(i) for i, (key, *_ ) in enumerate(cols, start=1)}
    fil_col, reg_col = col_letter["FILIERE"], col_letter["REGION"]
    fil_range = f"Données!${fil_col}$2:${fil_col}${last_row}"
    reg_range = f"Données!${reg_col}$2:${reg_col}${last_row}"

    if has_status:
        stat_col = col_letter["STATUT_DOSSIER"]
        stat_range = f"Données!${stat_col}$2:${stat_col}${last_row}"
        full_stat_range = f"{stat_col}2:{stat_col}{last_row}"
        ws.conditional_formatting.add(
            full_stat_range,
            CellIsRule(operator="equal", formula=['"PAS OK"'], fill=RED_FILL, font=RED_FONT),
        )
        ws.conditional_formatting.add(
            full_stat_range,
            CellIsRule(operator="equal", formula=['"OK"'], fill=GREEN_FILL, font=GREEN_FONT),
        )

    filieres = [k for k, _ in counts_by(records, "FILIERE")]
    regions = [k for k, _ in counts_by(records, "REGION")]

    # --- Par Filière / Par Région -------------------------------------------
    def build_breakdown_sheet(name, field_label, values, data_range):
        ws_b = wb.create_sheet(name)
        ws_b["A1"] = f"Répartition des candidats par {field_label}"
        ws_b["A1"].font = TITLE_FONT
        ws_b["A2"] = f"Total candidats : {n}"
        ws_b["A2"].font = SUB_FONT

        headers = [field_label.capitalize(), "Nombre de candidats", "% du total"]
        if has_status:
            headers += ["Dont PAS OK", "% PAS OK"]
        start_row = 4
        style_header_row(ws_b, start_row, headers)
        ws_b.column_dimensions["A"].width = 32
        for col, width in zip("BCDE", (20, 12, 14, 12)):
            ws_b.column_dimensions[col].width = width

        for i, val in enumerate(values, start=1):
            r = start_row + i
            ws_b.cell(row=r, column=1, value=val).font = NORMAL_FONT
            ws_b.cell(row=r, column=1).border = BORDER
            b = ws_b.cell(row=r, column=2, value=f'=COUNTIF({data_range},A{r})')
            b.font, b.border, b.alignment = NORMAL_FONT, BORDER, Alignment(horizontal="center")
            c = ws_b.cell(row=r, column=3, value=f'=B{r}/{n}')
            c.number_format = "0.0%"
            c.font, c.border, c.alignment = NORMAL_FONT, BORDER, Alignment(horizontal="center")
            if has_status:
                d = ws_b.cell(row=r, column=4, value=f'=COUNTIFS({data_range},A{r},{stat_range},"PAS OK")')
                d.font, d.border, d.alignment = NORMAL_FONT, BORDER, Alignment(horizontal="center")
                e = ws_b.cell(row=r, column=5, value=f'=IF(B{r}=0,0,D{r}/B{r})')
                e.number_format = "0.0%"
                e.font, e.border, e.alignment = NORMAL_FONT, BORDER, Alignment(horizontal="center")

        total_row = start_row + len(values) + 1
        last_data_row = start_row + len(values)
        ws_b.cell(row=total_row, column=1, value="TOTAL").font = BOLD_FONT
        ws_b.cell(row=total_row, column=2, value=f"=SUM(B{start_row+1}:B{last_data_row})").font = BOLD_FONT
        pc = ws_b.cell(row=total_row, column=3, value=f"=SUM(C{start_row+1}:C{last_data_row})")
        pc.font, pc.number_format = BOLD_FONT, "0.0%"
        if has_status:
            ws_b.cell(row=total_row, column=4, value=f"=SUM(D{start_row+1}:D{last_data_row})").font = BOLD_FONT
            pe = ws_b.cell(row=total_row, column=5, value=f"=IF(B{total_row}=0,0,D{total_row}/B{total_row})")
            pe.font, pe.number_format = BOLD_FONT, "0.0%"
        for col in range(1, len(headers) + 1):
            ws_b.cell(row=total_row, column=col).fill = ACCENT_FILL
            ws_b.cell(row=total_row, column=col).border = BORDER
        return start_row, total_row

    fil_start, _ = build_breakdown_sheet("Par Filière", "filière", filieres, fil_range)
    reg_start, reg_total = build_breakdown_sheet("Par Région", "région d'origine", regions, reg_range)

    ws3 = wb["Par Région"]
    note = ws3.cell(row=reg_total + 2, column=1,
                     value=f"Note : les candidats étrangers sans région camerounaise renseignée sont "
                           f"regroupés sous « {FOREIGN_LABEL} ».")
    note.font = SUB_FONT
    note.alignment = Alignment(wrap_text=True)
    ws3.merge_cells(start_row=reg_total + 2, start_column=1, end_row=reg_total + 2, end_column=5)
    ws3.row_dimensions[reg_total + 2].height = 30

    # --- Filière x Région ----------------------------------------------------
    ws4 = wb.create_sheet("Filière x Région")
    ws4["A1"] = "Tableau croisé : Filière x Région d'origine"
    ws4["A1"].font = TITLE_FONT
    start_row4 = 3
    ws4.cell(row=start_row4, column=1, value="Filière").font = HEADER_FONT
    ws4.cell(row=start_row4, column=1).fill = HEADER_FILL
    ws4.cell(row=start_row4, column=1).border = BORDER
    ws4.column_dimensions["A"].width = 30

    for j, reg in enumerate(regions, start=2):
        c = ws4.cell(row=start_row4, column=j, value=reg)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", text_rotation=45, wrap_text=True)
        ws4.column_dimensions[get_column_letter(j)].width = 12
    total_col = len(regions) + 2
    tc = ws4.cell(row=start_row4, column=total_col, value="Total")
    tc.font, tc.fill, tc.border = HEADER_FONT, HEADER_FILL, BORDER
    ws4.row_dimensions[start_row4].height = 60

    for i, fil in enumerate(filieres, start=1):
        r = start_row4 + i
        ws4.cell(row=r, column=1, value=fil).font = NORMAL_FONT
        ws4.cell(row=r, column=1).border = BORDER
        for j, reg in enumerate(regions, start=2):
            cl = get_column_letter(j)
            f = f'=COUNTIFS({fil_range},$A{r},{reg_range},{cl}${start_row4})'
            cc = ws4.cell(row=r, column=j, value=f)
            cc.font, cc.border = NORMAL_FONT, BORDER
            cc.alignment = Alignment(horizontal="center")
        trange = f"B{r}:{get_column_letter(total_col-1)}{r}"
        tcell = ws4.cell(row=r, column=total_col, value=f"=SUM({trange})")
        tcell.font, tcell.border = BOLD_FONT, BORDER
        tcell.alignment = Alignment(horizontal="center")

    total_row4 = start_row4 + len(filieres) + 1
    ws4.cell(row=total_row4, column=1, value="TOTAL").font = BOLD_FONT
    ws4.cell(row=total_row4, column=1).fill = ACCENT_FILL
    for j in range(2, total_col + 1):
        cl = get_column_letter(j)
        colrange = f"{cl}{start_row4+1}:{cl}{start_row4+len(filieres)}"
        cc = ws4.cell(row=total_row4, column=j, value=f"=SUM({colrange})")
        cc.font, cc.fill, cc.border = BOLD_FONT, ACCENT_FILL, BORDER
        cc.alignment = Alignment(horizontal="center")
    ws4.freeze_panes = ws4.cell(row=start_row4 + 1, column=2).coordinate

    # --- Filière x Statut / Région x Statut (only if status data given) -----
    def build_status_crosstab(name, field_label, values, data_range):
        ws_c = wb.create_sheet(name)
        ws_c["A1"] = f"Tableau croisé : {field_label} x Statut du dossier"
        ws_c["A1"].font = TITLE_FONT
        ws_c["A2"] = "« PAS OK » = statut différent de « en cours de traitement » (rapport de vérification)"
        ws_c["A2"].font = SUB_FONT
        ws_c.merge_cells("A2:E2")

        start = 4
        style_header_row(ws_c, start, [field_label, "OK", "PAS OK", "Total", "% PAS OK"])
        ws_c.column_dimensions["A"].width = 32
        for col in "BCDE":
            ws_c.column_dimensions[col].width = 12

        key_field = "FILIERE" if field_label == "Filière" else "REGION"
        pas_ok_counts = Counter(rec[key_field] for rec in records if rec["STATUT_DOSSIER"] == "PAS OK")
        ordered = sorted(values, key=lambda v: (-pas_ok_counts.get(v, 0), v))

        for i, val in enumerate(ordered, start=1):
            r = start + i
            ws_c.cell(row=r, column=1, value=val).font = NORMAL_FONT
            ws_c.cell(row=r, column=1).border = BORDER
            ok_c = ws_c.cell(row=r, column=2, value=f'=COUNTIFS({data_range},$A{r},{stat_range},"OK")')
            ok_c.font, ok_c.border = NORMAL_FONT, BORDER
            ok_c.alignment = Alignment(horizontal="center")
            nok_c = ws_c.cell(row=r, column=3, value=f'=COUNTIFS({data_range},$A{r},{stat_range},"PAS OK")')
            nok_c.font, nok_c.border = NORMAL_FONT, BORDER
            nok_c.alignment = Alignment(horizontal="center")
            tot_c = ws_c.cell(row=r, column=4, value=f'=COUNTIF({data_range},$A{r})')
            tot_c.font, tot_c.border = BOLD_FONT, BORDER
            tot_c.alignment = Alignment(horizontal="center")
            pct_c = ws_c.cell(row=r, column=5, value=f'=IF(D{r}=0,0,C{r}/D{r})')
            pct_c.number_format = "0.0%"
            pct_c.font, pct_c.border = NORMAL_FONT, BORDER
            pct_c.alignment = Alignment(horizontal="center")

        total_row = start + len(ordered) + 1
        last_data_row = start + len(ordered)
        ws_c.cell(row=total_row, column=1, value="TOTAL").font = BOLD_FONT
        for col, letter in zip((2, 3, 4), "BCD"):
            ws_c.cell(row=total_row, column=col,
                      value=f"=SUM({letter}{start+1}:{letter}{last_data_row})").font = BOLD_FONT
        pe = ws_c.cell(row=total_row, column=5, value=f"=IF(D{total_row}=0,0,C{total_row}/D{total_row})")
        pe.font, pe.number_format = BOLD_FONT, "0.0%"
        for col in range(1, 6):
            ws_c.cell(row=total_row, column=col).fill = ACCENT_FILL
            ws_c.cell(row=total_row, column=col).border = BORDER
        ws_c.conditional_formatting.add(
            f"C{start+1}:C{last_data_row}",
            CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL),
        )

    if has_status:
        build_status_crosstab("Filière x Statut", "Filière", filieres, fil_range)
        build_status_crosstab("Région x Statut", "Région d'origine", regions, reg_range)

    # --- Récapitulatif ---------------------------------------------------
    ws5 = wb.create_sheet("Récapitulatif")
    wb.move_sheet("Récapitulatif", offset=-(len(wb.sheetnames) - 1))

    title = exam_title or "RÉCAPITULATIF DES INSCRIPTIONS"
    ws5["B2"] = f"RÉCAPITULATIF — {title.upper()}"
    ws5["B2"].font = TITLE_FONT
    ws5.merge_cells("B2:F2")
    matricules_sorted = sorted(rec["MATRICULE"] for rec in records)
    subtitle = f"Fiches d'inscription {matricules_sorted[0]} à {matricules_sorted[-1]} ({n} dossiers)"
    if has_status:
        subtitle += " — croisé avec le rapport de vérification des dossiers"
    ws5["B3"] = subtitle
    ws5["B3"].font = SUB_FONT
    ws5.merge_cells("B3:F3")

    kpis = [
        ("Nombre total de candidats (fiches PDF)", f"=COUNTA(Données!A2:A{last_row})"),
        ("Nombre de filières distinctes", f"=SUMPRODUCT(1/COUNTIF({fil_range},{fil_range}))"),
        ("Nombre de régions/origines distinctes", f"=SUMPRODUCT(1/COUNTIF({reg_range},{reg_range}))"),
        ("Candidats de sexe féminin", f'=COUNTIF(Données!{col_letter["SEXE"]}2:{col_letter["SEXE"]}{last_row},"FEMININ")'),
        ("Candidats de sexe masculin", f'=COUNTIF(Données!{col_letter["SEXE"]}2:{col_letter["SEXE"]}{last_row},"MASCULIN")'),
        ("Candidats de nationalité étrangère", f'=COUNTIF({reg_range},"{FOREIGN_LABEL}")'),
    ]
    if has_status:
        kpis += [
            ("Dossiers OK (en cours de traitement)", f'=COUNTIF({stat_range},"OK")'),
            ("Dossiers PAS OK", f'=COUNTIF({stat_range},"PAS OK")'),
        ]

    r0 = 5
    for i, h in enumerate(["Indicateur", "Valeur"], start=2):
        c = ws5.cell(row=r0, column=i, value=h)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws5.column_dimensions["B"].width = 40
    ws5.column_dimensions["C"].width = 14
    pas_ok_row = None
    total_row_kpi = None
    for i, (label, formula) in enumerate(kpis, start=1):
        r = r0 + i
        ws5.cell(row=r, column=2, value=label).font = NORMAL_FONT
        ws5.cell(row=r, column=2).border = BORDER
        v = ws5.cell(row=r, column=3, value=formula)
        v.font, v.border, v.alignment = BOLD_FONT, BORDER, Alignment(horizontal="center")
        if label.startswith("Nombre total"):
            total_row_kpi = r
        if label == "Dossiers PAS OK":
            pas_ok_row = r

    if has_status:
        r = r0 + len(kpis) + 1
        ws5.cell(row=r, column=2, value="% de dossiers PAS OK").font = NORMAL_FONT
        ws5.cell(row=r, column=2).border = BORDER
        v = ws5.cell(row=r, column=3, value=f"=C{pas_ok_row}/C{total_row_kpi}")
        v.font, v.border, v.alignment = BOLD_FONT, BORDER, Alignment(horizontal="center")
        v.number_format = "0.0%"
        kpi_end = r
    else:
        kpi_end = r0 + len(kpis)

    def top5_block(top_row, field_sheet, label):
        ws5.cell(row=top_row, column=2, value=f"Top 5 des {label} les plus demandées"
                 if field_sheet == "Par Filière" else f"Top 5 des {label}").font = Font(
            name=FONT_NAME, bold=True, size=12, color="1F4E78")
        ws5.merge_cells(start_row=top_row, start_column=2, end_row=top_row, end_column=5)
        hr = top_row + 1
        headers = ["Filière" if field_sheet == "Par Filière" else "Région", "Nombre", "% du total"]
        if has_status:
            headers.append("Dont PAS OK")
        for i, h in enumerate(headers, start=2):
            c = ws5.cell(row=hr, column=i, value=h)
            c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
            c.alignment = Alignment(horizontal="center")
        values = filieres if field_sheet == "Par Filière" else regions
        src_start = fil_start if field_sheet == "Par Filière" else reg_start
        for i in range(1, min(5, len(values)) + 1):
            r = hr + i
            ws5.cell(row=r, column=2, value=f"='{field_sheet}'!A{src_start+i}").font = NORMAL_FONT
            ws5.cell(row=r, column=2).border = BORDER
            ws5.cell(row=r, column=3, value=f"='{field_sheet}'!B{src_start+i}").font = NORMAL_FONT
            ws5.cell(row=r, column=3).border = BORDER
            ws5.cell(row=r, column=3).alignment = Alignment(horizontal="center")
            pc = ws5.cell(row=r, column=4, value=f"='{field_sheet}'!C{src_start+i}")
            pc.number_format = "0.0%"
            pc.font, pc.border = NORMAL_FONT, BORDER
            pc.alignment = Alignment(horizontal="center")
            if has_status:
                dc = ws5.cell(row=r, column=5, value=f"='{field_sheet}'!D{src_start+i}")
                dc.font, dc.border = NORMAL_FONT, BORDER
                dc.alignment = Alignment(horizontal="center")
        return hr + 6

    next_row = top5_block(kpi_end + 3, "Par Filière", "filières")
    next_row = top5_block(next_row + 2, "Par Région", "régions d'origine")

    notes = [
        f"Sources : {n} fiches d'inscription PDF (matricules {matricules_sorted[0]} à {matricules_sorted[-1]}).",
    ]
    if orphan_matricules:
        notes.append(
            f"{len(orphan_matricules)} matricule(s) du rapport de statut n'ont pas de fiche PDF "
            f"correspondante et n'apparaissent donc pas dans l'onglet Données : "
            + ", ".join(orphan_matricules) + "."
        )
    if FILIERE_ALIASES:
        aliases_txt = "; ".join(f"« {k} » → « {v} »" for k, v in FILIERE_ALIASES.items())
        notes.append(f"Filières harmonisées pour le classement : {aliases_txt}.")
    if has_status:
        notes.append("« PAS OK » signifie : statut différent de « en cours de traitement ».")

    note_row = next_row + 2
    note_cell = ws5.cell(row=note_row, column=2, value=" ".join(notes))
    note_cell.font = SUB_FONT
    note_cell.alignment = Alignment(wrap_text=True)
    ws5.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=6)
    ws5.row_dimensions[note_row].height = 60

    wb.active = 0
    wb.save(out_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--zip", required=True, help="ZIP archive of fiche d'inscription PDFs")
    parser.add_argument("--rapport", help="Optional .docx dossier-status report")
    parser.add_argument("--out", required=True, help="Output .xlsx path")
    args = parser.parse_args()

    records, skipped, exam_title = extract_records_from_zip(args.zip)
    if not records:
        sys.exit("Aucune fiche exploitable trouvée dans le ZIP.")

    print(f"{len(records)} fiche(s) extraite(s).")
    if skipped:
        print(f"{len(skipped)} fichier(s) ignoré(s) :")
        for name, reason in skipped:
            print(f"  - {name}: {reason}")
    for rec in records:
        if rec.get("_MISSING_FIELDS"):
            print(f"  ! {rec['MATRICULE']}: champs manquants {rec['_MISSING_FIELDS']}")

    orphan_matricules = []
    has_status = False
    if args.rapport:
        status_map = parse_status_report(args.rapport)
        orphan_matricules = merge_status(records, status_map)
        has_status = True
        print(f"Statut croisé pour {len(status_map)} matricule(s) du rapport "
              f"({len(orphan_matricules)} sans fiche correspondante).")

    build_workbook(records, args.out, exam_title, orphan_matricules, has_status)
    print(f"Classeur généré : {args.out}")


if __name__ == "__main__":
    main()
